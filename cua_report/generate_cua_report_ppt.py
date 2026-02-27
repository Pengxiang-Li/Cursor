from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Literal

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


SlideKind = Literal["title", "bullets", "table"]


@dataclass(frozen=True)
class Bullet:
    level: int
    text: str


@dataclass(frozen=True)
class SlideSpec:
    kind: SlideKind
    title: str
    subtitle: str | None = None
    bullets: list[Bullet] | None = None
    table_columns: list[str] | None = None
    table_rows: list[list[str]] | None = None
    notes: list[str] | None = None


DEFAULT_FONT = "Noto Sans CJK SC"


def _set_run_font(run, *, size_pt: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = DEFAULT_FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_notes(slide, notes: list[str] | None) -> None:
    if not notes:
        return
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    tf.clear()
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = "讲稿备注（建议口径）"
    p0.level = 0
    for run in p0.runs:
        _set_run_font(run, size_pt=14, bold=True)
    for line in notes:
        p = tf.add_paragraph()
        p.text = line
        p.level = 1
        for run in p.runs:
            _set_run_font(run, size_pt=12)


def _add_title_slide(prs: Presentation, spec: SlideSpec) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # title
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title_tf = title.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    r = p.add_run()
    r.text = spec.title
    p.alignment = PP_ALIGN.LEFT
    _set_run_font(r, size_pt=40, bold=True)

    subtitle_tf = subtitle.text_frame
    subtitle_tf.clear()
    p2 = subtitle_tf.paragraphs[0]
    r2 = p2.add_run()
    r2.text = spec.subtitle or ""
    _set_run_font(r2, size_pt=18)

    _add_notes(slide, spec.notes)


def _add_bullet_slide(prs: Presentation, spec: SlideSpec) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    title = slide.shapes.title
    body = slide.shapes.placeholders[1]

    title_tf = title.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    r = p.add_run()
    r.text = spec.title
    _set_run_font(r, size_pt=32, bold=True)

    tf = body.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE

    bullets = spec.bullets or []
    if not bullets:
        p0 = tf.paragraphs[0]
        p0.text = ""
    else:
        first = True
        for b in bullets:
            if first:
                p0 = tf.paragraphs[0]
                p0.text = b.text
                p0.level = max(0, b.level)
                for run in p0.runs:
                    _set_run_font(run, size_pt=18)
                first = False
            else:
                pN = tf.add_paragraph()
                pN.text = b.text
                pN.level = max(0, b.level)
                for run in pN.runs:
                    _set_run_font(run, size_pt=18)

    _add_notes(slide, spec.notes)


def _add_table_slide(prs: Presentation, spec: SlideSpec) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    title = slide.shapes.title

    title_tf = title.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    r = p.add_run()
    r.text = spec.title
    _set_run_font(r, size_pt=30, bold=True)

    cols = spec.table_columns or []
    rows = spec.table_rows or []
    n_rows = max(1, len(rows) + 1)
    n_cols = max(1, len(cols))

    left = Inches(0.6)
    top = Inches(1.6)
    width = Inches(12.2)
    height = Inches(5.6)
    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
    table = table_shape.table

    # header
    for j, col in enumerate(cols):
        cell = table.cell(0, j)
        cell.text = col
        for p in cell.text_frame.paragraphs:
            for run in p.runs:
                _set_run_font(run, size_pt=14, bold=True, color=RGBColor(255, 255, 255))
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(31, 78, 121)

    # body
    for i, row in enumerate(rows, start=1):
        for j in range(n_cols):
            cell = table.cell(i, j)
            cell.text = row[j] if j < len(row) else ""
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    _set_run_font(run, size_pt=12)

    _add_notes(slide, spec.notes)


def _date_str() -> str:
    return _dt.date.today().isoformat()


def build_slides() -> list[SlideSpec]:
    today = _date_str()
    return [
        SlideSpec(
            kind="title",
            title="CUA 调研汇报：从现有工作抽象出方法论",
            subtitle=f"聚焦：数据飞轮 / Grounding / 轨迹生产 / RL+Verifier / 动作空间统一（{today}）",
            notes=[
                "开场 30 秒：这次不讲单篇论文细节，讲跨工作可复用的方法论（recipe）。",
                "听众默认懂训练与 RL；我只强调那些“落到工程与数据闭环才会出现”的关键点。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="目标与输出",
            bullets=[
                Bullet(0, "目标：分析现有 SOTA 工作，抽象出可执行的方法论（recipe）"),
                Bullet(0, "输出 1：一条端到端训练路线图（Grounding → SFT → RL → 数据飞轮）"),
                Bullet(0, "输出 2：关键设计分歧与 trade-off（Vision-only vs A11y、动作空间、Verifier）"),
                Bullet(0, "输出 3：可落地的工程与数据管线拆解（环境并发、重置、路由、质检）"),
            ],
            notes=[
                "把“论文点”变成“我们明天能搭建”的组件：动作空间、数据 schema、奖励与环境。",
                "所有结论尽量用多篇工作共同出现的信号来支撑，避免单点经验。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="问题定义：CUA（Computer-Use Agent）系统栈",
            bullets=[
                Bullet(0, "输入：截图（可选 A11y/DOM/UI tree）+ 历史（动作摘要/记忆）+ 任务指令"),
                Bullet(0, "输出：动作（坐标 primitives）或代码（API-GUI paradigm）或工具调用（MCP）"),
                Bullet(0, "闭环：执行 → 观测 → 规划/反思 → 再执行；外部 Verifier 提供成败信号"),
                Bullet(0, "核心指标：Task Success Rate + Efficiency（步数/耗时）+ 鲁棒性（异常/不可行）"),
            ],
            notes=[
                "先把系统边界画清：这是多模态序列决策 + 真实环境反馈（往往稀疏且昂贵）。",
                "后续所有方法论都在解决三件事：可学习的观测、可控的动作、可扩展的奖励/验证。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="跨工作共识 1：数据闭环（Data Flywheel）> 模型架构",
            bullets=[
                Bullet(0, "UI-TARS-2：Data Flywheel + 数据路由（高质→SFT，低质→CT）"),
                Bullet(0, "Computer-RL：RL 触顶后用 Entropulse（成功 rollouts 再 SFT）恢复熵与探索"),
                Bullet(0, "Mano：在线 RL 成功轨迹回流到 SFT；“成功但有错”→LLM 草稿→人修→再注入"),
                Bullet(0, "EvoCUA：5K 轨迹多轮 RFT + DPO 即可上分，强调迭代质量与 scaling 基建"),
            ],
            notes=[
                "把这页当成总论：几乎所有 SOTA 都在“数据生产—训练—再生产”上卷效率，而不是换 backbone。",
                "后面每个模块都要回答：它如何进入闭环、如何被质检、如何被路由到下一轮训练。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="跨工作共识 2：三大瓶颈（按影响排序）",
            bullets=[
                Bullet(0, "Grounding：桌面元素密/小/多窗口；ScreenSpot-Pro 仍很难（材料中提到 18.9%）"),
                Bullet(0, "RL Scaling：性能跃迁依赖 RL，但瓶颈在 环境并发 + 快速重置 + Verifier 质量"),
                Bullet(0, "动作空间割裂：Mobile/Web/Desktop 不一致；梯度冲突/迁移困难（MRPO、插值等都在缓解）"),
            ],
            notes=[
                "这三件事决定“能不能做大”：Grounding 决定能不能点对；Verifier 决定 RL 上限；动作空间决定跨平台与工具化能力。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="动作空间：从坐标 JSON 到 API-GUI Paradigm（Computer-RL）",
            bullets=[
                Bullet(0, "现状：click/type/drag/scroll 等 primitives 在不同平台/论文里 schema 不一致"),
                Bullet(0, "Computer-RL：直接输出 Python 代码块，统一 GUI primitive + app-specific API"),
                Bullet(0, "Mano：Thought 与 Action 之间插入 Action Desp（一句话摘要）→ 显著提分（+2.8）"),
                Bullet(0, "Mobile-Agent-v3.5：训练-推理对齐用 Token-ID Transport，避免 logprob 错位"),
            ],
            notes=[
                "给资深听众的点：动作空间不是 UI 细节，而是 credit assignment 与可组合性（循环/条件/函数）。",
                "如果要做跨平台统一，代码动作空间是目前最“端到端对齐编程能力”的方案。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="数据版图：Grounding / Trajectory / Cognition（思考-记忆-世界模型）",
            bullets=[
                Bullet(0, "Grounding：Text/Icon → Point/BBox；含 infeasible 负样本（拒答/纠错）"),
                Bullet(0, "Trajectory：多步交互 (screenshot, thought, action)；可混 A11y/DOM 元数据"),
                Bullet(0, "Cognition：Observation/Memory/Reflection/Progress → Thought → Action（Mobile-Agent v3.5）"),
                Bullet(0, "World Modeling：预测“动作后界面变化”作为监督信号（Mobile-Agent v3.5）"),
            ],
            notes=[
                "强调一个组织方式：把数据按“训练目标”分桶，而不是按“来源”分桶，便于配比与路由。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="Grounding 数据工程：高难样本 + 负样本 + 轨迹挖掘",
            bullets=[
                Bullet(0, "AgentCPM：12M grounding 样本（并混 50% 通用多模态做正则）"),
                Bullet(0, "Mobile-Agent-v3.5：Hard Grounding Synthesis（专业软件/多窗口/高分辨率）"),
                Bullet(0, "Mobile-Agent-v3.5：Infeasible 负样本（截图×Query 随机组合 + 多模型共识过滤）"),
                Bullet(0, "Mobile-Agent-v3.5：从海量探索轨迹用 Critic 挖掘/清洗 grounding 对（低成本扩容）"),
            ],
            notes=[
                "Grounding 的关键不只是数量，而是“错在哪里”：小控件、遮挡、多窗口、语义歧义。",
                "负样本在 CUA 很关键：能显著降低盲点点击与 hallucinated element。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="Trajectory 生产：三条路线（规模×质量）",
            bullets=[
                Bullet(0, "自动化探索（Mobile-Agent-v3.5）：人工建 DAG + checkpoint；失败回滚+重写指令→干净轨迹"),
                Bullet(0, "虚拟环境（Mobile-Agent-v3.5）：Web 渲染可控 app，绕开验证码/噪声，长轨迹高效生成"),
                Bullet(0, "自动化采集（Mano Explorer）：元素提取（Web 插件/桌面 A11y+OmniParse）+ DFS 深度≤10"),
                Bullet(0, "人工标注：只兜底长尾困难任务；可做 in-policy（UI-TARS-2：accept/override）"),
            ],
            notes=[
                "这里的主线是：用系统结构（DAG/虚拟环境/质量评估）把‘轨迹质量’工程化，而不是堆人力。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="任务生成与难度分层：让数据更像推理时的分布",
            bullets=[
                Bullet(0, "UI-TARS-2：用 LLM 生成 Multi-Condition Obfuscation / Multi-Hop 条件指令"),
                Bullet(0, "Computer-RL：按难度分层采集（Easy/Medium/Hard）+ Model Pool 教师多样性"),
                Bullet(0, "EvoCUA：同一 Query 多条正确轨迹 > 更多 Query（覆盖 policy space）"),
                Bullet(0, "关键步骤（EvoCUA）：普通步骤 vs 关键步骤 upsampling；关键步可用 offline DPO 专训"),
            ],
            notes=[
                "把任务难度分层，是为了把采样预算用在能提升 policy 的地方（hard tasks / high-entropy steps）。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="训练范式总览：一条可复现的 recipe",
            bullets=[
                Bullet(0, "（可选）CT/Pretrain：UI 感知 + OCR + grounding + 工具语义"),
                Bullet(0, "SFT：带 CoT 的多步轨迹对齐（历史帧/摘要/记忆）"),
                Bullet(0, "Online RL：GRPO/PPO 触发性能跃迁（稀疏 reward + 强 Verifier）"),
                Bullet(0, "Entropulse：用成功 rollouts SFT 一轮恢复熵，再继续 RL（Computer-RL）"),
                Bullet(0, "Offline DPO：对关键步骤或常错模式做样本效率更高的偏好学习（EvoCUA/Mano）"),
                Bullet(0, "跨平台融合：交替优化（MRPO）或专项模型→参数插值（UI-TARS-2）"),
            ],
            notes=[
                "这一页是后面细节的索引：你每做一步，都要说明它在闭环里怎么回流与路由。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="SFT 细节：输入组织（历史帧）与输出结构（摘要/CoT）",
            bullets=[
                Bullet(0, "历史帧策略（Mano）：保留前 2 帧历史图像 + 全部历史文本摘要 → 性能/开销最优"),
                Bullet(0, "Action Desp（Mano）：在 Thought 与 Action 之间加一句话摘要，强约束下一步"),
                Bullet(0, "统一 CoT 合成（Mobile-Agent-v3.5）：Observation/Memory/Reflection/Progress → Thought → Action"),
                Bullet(0, "世界模型监督（Mobile-Agent-v3.5）：预测 action 后界面变化，提升前瞻性"),
            ],
            notes=[
                "对资深听众：这本质是在做‘信息瓶颈设计’——哪些信息应该进 context、哪些应该被压缩成摘要。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="RL 细节：算法选择、奖励设计、以及训练-推理一致性工程",
            bullets=[
                Bullet(0, "Step-wise GRPO（Computer-RL/DART-GUI）：不需要 Value Network，显存友好；group 内 z-score advantage"),
                Bullet(0, "PPO（UI-TARS-2）：更稳定，但需 Value 预训练；长序列用 Decoupled/Length-Adaptive GAE"),
                Bullet(0, "MRPO（Mobile-Agent-v3.5）：交替多平台优化，缓解多设备梯度 tug-of-war"),
                Bullet(0, "Reward：Rule-based verifier（优先）→ screenshot-diff+OCR → LLM-as-judge → Generative ORM（UI-TARS-2）"),
                Bullet(0, "一致性工程（Mobile-Agent-v3.5）：Token-ID Transport，避免 logprob 计算错位"),
                Bullet(0, "熵坍塌：RL 后期 mode collapse → Entropulse（Computer-RL）恢复探索"),
            ],
            notes=[
                "强调一个‘必须做’：reward/verification 的分层；否则 online RL 很难 scale 到桌面任务。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="Verifier 体系：决定 RL 上限的“隐形模型”",
            bullets=[
                Bullet(0, "Tier 1 Rule-based：文件系统/窗口状态/文档结构解析/配置读取（最可靠）"),
                Bullet(0, "Tier 2 Screenshot-diff：关键区域比对 + OCR/regex（中等可扩展）"),
                Bullet(0, "Tier 3 LLM-as-Judge：灵活但噪声大；可输出标量评分与失败原因"),
                Bullet(0, "Tier 4 Generative ORM（UI-TARS-2）：自训练 outcome reward model，统一开放任务打分"),
            ],
            notes=[
                "讲清楚优先级：先把 P0 应用做成 Tier1 verifier（可编程/可解析），再用 Tier3 覆盖长尾。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="环境与系统架构：让 Online RL 变成可规模化的工程",
            bullets=[
                Bullet(0, "并发需求：EvoCUA 提到 2000–4000 环境；Computer-RL/相关工作也依赖数千并发"),
                Bullet(0, "核心工程点：快照池 + 快速重置（<10s）+ 轨迹为最小调度单元（rollout-wise）"),
                Bullet(0, "解耦异步架构（DART-GUI）：Env Cluster / Rollout Service / Data Manager / Trainer"),
                Bullet(0, "near on-policy：每条轨迹只用一次 + 渐进式同步，减少分布漂移与全局 barrier"),
            ],
            notes=[
                "这页的讲法偏工程：强调 RL 算法不是瓶颈，瓶颈是环境调度与验证吞吐。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="观测输入：Vision-only vs A11y Tree（以及如何混合）",
            bullets=[
                Bullet(0, "A11y 的价值：可给结构化元素/role/name/坐标；更易做 grounding 与解释"),
                Bullet(0, "A11y 的代价：token 成本高、跨 OS 覆盖不一致、Electron/浏览器需额外 flag"),
                Bullet(0, "Computer-RL：明确移除 A11y tree，减少 token 且避免冗余（纯视觉仍可 SOTA）"),
                Bullet(0, "建议：训练时混合（有/无 A11y），推理时按任务/应用选择；数据 schema 需兼容"),
            ],
            notes=[
                "把它讲成一个工程策略：A11y 不是‘要不要’，而是‘什么时候用’（gating）。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="桌面应用选择：可程序化 + 易验证优先（来自环境选型分析）",
            bullets=[
                Bullet(0, "跨平台 P0：文件管理器 / 浏览器 / 终端 / Office 套件 / VS Code / 系统设置 / 文本编辑器"),
                Bullet(0, "最佳 code 自动化：Browser=Playwright；Office=Win COM / Linux UNO / macOS AppleScript+UNO"),
                Bullet(0, "A11y 基建：Linux AT-SPI（Wayland 限制键鼠注入）、Windows UIA、macOS AX（需授权）"),
                Bullet(0, "Verifier 难度：文件/终端/文本（低）→ 浏览器/Office/设置（中）→ 图像编辑/邮件（高）"),
            ],
            notes=[
                "强调一个决策原则：先做能写 verifier 的应用（否则 RL 没法闭环）。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="方法论：8 步搭建你的 CUA 基座（把论文共识变成执行清单）",
            bullets=[
                Bullet(0, "1) 定义统一动作空间：优先考虑代码动作空间（API-GUI paradigm）+ primitives 兜底"),
                Bullet(0, "2) 选 P0 应用与任务域：可程序化/可验证/高频；先覆盖文件/浏览器/Office/终端"),
                Bullet(0, "3) 先建 Verifier（Tier1→Tier3）与快照重置，再谈 Online RL"),
                Bullet(0, "4) Grounding 先行：高难合成 + 负样本 + 轨迹挖掘，持续提升点击正确率"),
                Bullet(0, "5) 冷启动 SFT：高质量短轨迹 + 2 帧历史图像 + 全历史摘要 + Action Desp"),
                Bullet(0, "6) 启动数据飞轮：rollout→验证→路由（成功→SFT；失败/低质→CT 或修正）"),
                Bullet(0, "7) Online RL + 熵管理：GRPO/PPO 触顶后用 Entropulse 再继续"),
                Bullet(0, "8) 跨平台融合：交替优化（MRPO）或专项模型→参数插值（更稳）"),
            ],
            notes=[
                "这页作为你的‘总结落点’：听众带走的是这 8 步，而不是某个论文 trick。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="常见 Failure Modes → 对策（从材料中抽象）",
            bullets=[
                Bullet(0, "Grounding 错误：小控件/遮挡/多窗口 → hard synthesis + 负样本 + critic 挖掘"),
                Bullet(0, "RL 熵坍塌：mode collapse/过度自信 → Entropulse（成功 rollouts SFT）+ 熵监控"),
                Bullet(0, "跨设备梯度冲突：batch 混训 tug-of-war → MRPO 交替训练 / 专项模型插值融合"),
                Bullet(0, "训练-推理不一致：logprob 错位 → Token-ID Transport（回传 token ids）"),
                Bullet(0, "Verifier 缺口：开放任务难验证 → 分层 verifier + Generative ORM"),
                Bullet(0, "系统约束：Wayland 键鼠注入、macOS 授权 → 工具链预配置 + 纯视觉兜底"),
            ],
            notes=[
                "讲法：每条 failure 都要落到‘能否自动化发现/自动化修复/进入闭环’。",
            ],
        ),
        SlideSpec(
            kind="table",
            title="开源数据与评测地形图（如何组合使用）",
            table_columns=["模块", "优先资源（材料覆盖）", "用途/注意点"],
            table_rows=[
                ["Grounding 预训练", "OS-ATLAS / GroundCUA / Jedi / ScreenSpot(-Pro)", "先把点击/定位做对；加入 infeasible 负样本提升拒答/纠错"],
                ["轨迹 SFT", "OpenCUA/AgentNet / GUI-360° / ScaleCUA / OmniACT", "统一 schema；OpenCUA 与基座 pattern 差异需对齐（材料提到 RoPE 等）"],
                ["端到端评测", "OSWorld / Windows Agent Arena / AssistGUI", "以 Task SR 为主；搭配效率指标与失败分析"],
                ["综合评测", "MMBench-GUI / UI-Vision", "分层看能力：理解→定位→动作→协作"],
                ["在线 RL", "自建任务 + 可编程 verifier（Computer-RL 提到 ~8000）", "瓶颈在环境并发与 verifier 吞吐；先做 P0 应用闭环"],
            ],
            notes=[
                "这页帮助大家把资源拼起来：不要把 OSWorld 当训练集；把它当验证与 failure mining。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="结论与下一步实验（面向方法论验证）",
            bullets=[
                Bullet(0, "最重要的不是 backbone，而是：动作空间 + Verifier + 数据飞轮的迭代效率"),
                Bullet(0, "建议的 ablation：A11y on/off、动作空间（JSON vs code）、Entropulse 有无、关键步骤 DPO"),
                Bullet(0, "关键未解：跨平台统一动作空间、自动 verifier 生成/自举、开放任务的可验证化"),
                Bullet(0, "可选方向：做一个可扩展的开源轨迹/环境/验证基建（数据管线 paper 也成立）"),
            ],
            notes=[
                "收尾口径：如果只能投一个月资源，先把 P0 应用的 verifier+RL 闭环跑通，再谈泛化。",
            ],
        ),
        SlideSpec(
            kind="bullets",
            title="References（材料覆盖的关键工作）",
            bullets=[
                Bullet(0, "Mobile-Agent-v3.5 / Mano Tech Report / UI-TARS-2 / EvoCUA / AgentCPM"),
                Bullet(0, "Computer-RL（AutoGLM-OS）/ DART-GUI / OpenCUA/AgentNet / OSWorld"),
                Bullet(0, "OS-ATLAS / GroundCUA / ScreenSpot-Pro / GUI-360° / ScaleCUA / MMBench-GUI"),
            ],
            notes=[
                "如果被问细节：把讨论拉回到‘闭环如何搭建’与‘哪一步是瓶颈’。",
            ],
        ),
    ]


def build_ppt(out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)

    for spec in build_slides():
        if spec.kind == "title":
            _add_title_slide(prs, spec)
        elif spec.kind == "bullets":
            _add_bullet_slide(prs, spec)
        elif spec.kind == "table":
            _add_table_slide(prs, spec)
        else:
            raise ValueError(f"Unknown slide kind: {spec.kind}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def main() -> None:
    out = Path(__file__).resolve().parent / "CUA_research_methodology.pptx"
    path = build_ppt(out)
    print(str(path))


if __name__ == "__main__":
    main()

